#!/usr/bin/env python3
"""MANUAL-TESTING v1.10 postcheck — the docx structural + content-marker
verification (the M54/M55 pattern: opens, sections, tables, h1/h2 counts,
all content markers) + the LibreOffice PDF text-verification round."""
import subprocess
import sys

DOCX = 'download/FiberOps-Manual-Testing-Guide.docx'

# ── 1. the docx structural check ─────────────────────────────────────
try:
    from docx import Document
except ImportError:
    subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'python-docx'], check=True)
    from docx import Document

doc = Document(DOCX)
sections = len(doc.sections)
tables = len(doc.tables)
h1 = sum(1 for p in doc.paragraphs if p.style is not None and p.style.name == 'Heading 1')
h2 = sum(1 for p in doc.paragraphs if p.style is not None and p.style.name == 'Heading 2')
full = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full += '\n' + cell.text

markers = [
    'AC-25', 'AC-26', 'AC-27', 'AC-28', 'Version 1.10', '359ae5e', '1636',
    'route_smoke_m56', 'cheque', 'PDC', 'pay-pdc', 'post_cheque_clear',
    'post_cheque_bounce', 'chequeStatus', 'OVERDUE', 'M56', '79 files',
]
missing = [m for m in markers if m not in full]
print(f'docx: opens=True sections={sections} tables={tables} h1={h1} h2={h2}')
print(f'content markers: {len(markers) - len(missing)}/{len(markers)} present' + (f' MISSING={missing}' if missing else ''))
assert sections >= 3 and tables >= 50 and h1 >= 15 and h2 >= 50, 'structural regression'
assert not missing, f'missing markers: {missing}'
print('DOCX POSTCHECK OK')

# ── 2. the LibreOffice PDF text verification ─────────────────────────
subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', '/tmp/m56pdf', DOCX],
               capture_output=True, check=True, timeout=300)
pdf = '/tmp/m56pdf/FiberOps-Manual-Testing-Guide.pdf'
try:
    import pypdf
except ImportError:
    subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'pypdf'], check=True)
    import pypdf

reader = pypdf.PdfReader(pdf)
pages = len(reader.pages)
text = '\n'.join((p.extract_text() or '') for p in reader.pages)
pdf_markers = [
    'AC-25', 'AC-26', 'AC-27', 'AC-28', 'Version 1.10', '359ae5e', '1636',
    'route_smoke_m56', 'post_cheque_clear', 'post_cheque_bounce', 'PDC',
    'chequeStatus', 'OVERDUE', 'pay-pdc', 'M56',
]
pdf_missing = [m for m in pdf_markers if m not in text]
print(f'pdf: {pages} pages, markers {len(pdf_markers) - len(pdf_missing)}/{len(pdf_markers)} present' + (f' MISSING={pdf_missing}' if pdf_missing else ''))
assert not pdf_missing, f'pdf missing: {pdf_missing}'
print('PDF TEXT VERIFICATION OK')
