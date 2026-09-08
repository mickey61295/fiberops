#!/usr/bin/env python3
"""MANUAL-TESTING v1.9 postcheck — the docx structural + content-marker
verification (the M54 pattern: opens, sections, tables, h1/h2 counts, all
content markers) + the LibreOffice PDF text-verification round."""
import subprocess
import sys

DOCX = 'download/FiberOps-Manual-Testing-Guide.docx'

# ── 1. the docx structural check ─────────────────────────────────────
try:
    from docx import Document
except ImportError:
    subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'python-docx'], check=True)
    from docx import Document

doc = Document(DOCX)
sections = len(doc.sections)
tables = len(doc.tables)
h1 = sum(1 for p in doc.paragraphs if p.style is not None and p.style.name == 'Heading 1')
h2 = sum(1 for p in doc.paragraphs if p.style is not None and p.style.name == 'Heading 2')
full = '\n'.join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full += '\n' + cell.text

markers = [
    'HR-12', 'HR-13', 'HR-14', 'Version 1.9', '68c5034', '1618',
    'route_smoke_m55', 'shift wages', 'unassigned', 'hr-l06',
    'post_shift_wages', 'shift-wages', 'M55', '78 files',
]
missing = [m for m in markers if m not in full]
print(f'docx: opens=True sections={sections} tables={tables} h1={h1} h2={h2}')
print(f'content markers: {len(markers) - len(missing)}/{len(markers)} present' + (f' MISSING={missing}' if missing else ''))
assert sections >= 3 and tables >= 50 and h1 >= 15 and h2 >= 50, 'structural regression'
assert not missing, f'missing markers: {missing}'
print('DOCX POSTCHECK OK')

# ── 2. the LibreOffice PDF text verification ─────────────────────────
subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', '/tmp/m55pdf', DOCX],
               capture_output=True, check=True, timeout=300)
pdf = '/tmp/m55pdf/FiberOps-Manual-Testing-Guide.pdf'
try:
    import pypdf
except ImportError:
    subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'pypdf'], check=True)
    import pypdf

reader = pypdf.PdfReader(pdf)
pages = len(reader.pages)
pdftext = '\n'.join((p.extract_text() or '') for p in reader.pages)
pdf_markers = [
    'HR-12', 'HR-13', 'HR-14', 'Version 1.9', '68c5034', '1618',
    'route_smoke_m55', 'post_shift_wages', 'unassigned', 'hr-l06', 'M55',
]
pdf_missing = [m for m in pdf_markers if m not in pdftext]
print(f'pdf: pages={pages}')
print(f'pdf markers: {len(pdf_markers) - len(pdf_missing)}/{len(pdf_markers)} present' + (f' MISSING={pdf_missing}' if pdf_missing else ''))
assert pages >= 60, f'page count regression: {pages}'
assert not pdf_missing, f'pdf missing markers: {pdf_missing}'
print('PDF TEXT-VERIFY OK')
